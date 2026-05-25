from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_POST
from .models import Home, Student, Area
from .forms import HomeForm, StudentForm, AreaForm
from django.db.models import Q
from .utils import identify_and_swap
from decimal import Decimal
import openpyxl

@login_required
def member_dashboard(request):
    # If not a home member, go to main dashboard
    if request.user.role != 'home':
        return redirect('finance:dashboard')
    
    # If home member but no profile linked, just show empty state or error instead of redirecting back
    home = request.user.home_profile
    if not home:
        return render(request, 'homes/member_dashboard.html', {'no_profile': True})
    invoices = home.invoices.all().order_by('-month')
    total_due = sum((inv.balance for inv in invoices), Decimal('0.00'))
    students = home.students.all()
    
    context = {
        'home': home,
        'invoices': invoices,
        'total_due': total_due,
        'students': students,
    }
    return render(request, 'homes/member_dashboard.html', context)

@login_required
def home_list(request):
    query = request.GET.get('q')
    lang = request.GET.get('lang', 'ml')
    homes = Home.objects.all()
    if query:
        homes = homes.filter(
            Q(uid__icontains=query) |
            Q(name__icontains=query) |
            Q(name_en__icontains=query) |
            Q(house_name__icontains=query) |
            Q(contact_number__icontains=query) |
            Q(area__name__icontains=query) |
            Q(area__name_en__icontains=query)
        )
    return render(request, 'homes/home_list.html', {
        'homes': homes, 
        'query': query,
        'lang': lang
    })

@login_required
def home_create(request):
    if request.method == 'POST':
        form = HomeForm(request.POST)
        if form.is_valid():
            home = form.save()
            if request.POST.get('create_user_account'):
                create_user_for_home(home)
            return redirect('homes:home_list')
    else:
        form = HomeForm()
    return render(request, 'homes/home_form.html', {'form': form, 'title': 'Add New Home'})

@login_required
def home_update(request, pk):
    home = get_object_or_404(Home, pk=pk)
    if request.method == 'POST':
        form = HomeForm(request.POST, instance=home)
        if form.is_valid():
            home = form.save()
            if request.POST.get('create_user_account') and not hasattr(home, 'user_account'):
                create_user_for_home(home)
            return redirect('homes:home_list')
    else:
        form = HomeForm(instance=home)
    return render(request, 'homes/home_form.html', {'form': form, 'title': 'Update Home', 'edit': True, 'home_obj': home})

def create_user_for_home(home):
    from accounts.models import CustomUser
    import re
    # Create a simple username from the home name
    username = re.sub(r'[^a-zA-Z0-9]', '', home.name).lower()
    # Ensure unique username
    base_username = username
    counter = 1
    while CustomUser.objects.filter(username=username).exists():
        username = f"{base_username}{counter}"
        counter += 1
        
    user = CustomUser.objects.create(
        username=username,
        role='home',
        home_profile=home,
        first_name=home.name
    )
    user.set_password('madrasa123')
    user.save()
    return user

@login_required
def home_delete(request, pk):
    home = get_object_or_404(Home, pk=pk)
    if request.method == 'POST':
        home.delete()
        messages.success(request, f"Home '{home.name}' deleted successfully.")
        return redirect('homes:home_list')
    return render(request, 'homes/home_confirm_delete.html', {'home_obj': home})

@login_required
def home_bulk_delete(request):
    if request.method == 'POST':
        home_ids = request.POST.getlist('home_ids')
        if home_ids:
            count = Home.objects.filter(pk__in=home_ids).delete()[0]
            messages.success(request, f"Successfully deleted {count} homes.")
        else:
            messages.warning(request, "No homes selected for deletion.")
    return redirect('homes:home_list')

@login_required
def home_upload_excel(request):
    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, "Please upload a valid Excel file.")
            return redirect('homes:home_upload_excel')
            
        try:
            wb = openpyxl.load_workbook(excel_file)
            sheet = wb.active
            
            count = 0
            # Assuming row 1 is header
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0]: # name is required
                    continue
                    
                raw_name = str(row[0]).strip()
                raw_name_en = str(row[1]).strip() if row[1] else ''
                
                # Automatically identify and swap if necessary
                name, name_en = identify_and_swap(raw_name, raw_name_en)
                
                house_name = str(row[2]).strip() if row[2] else ''
                home_type = str(row[3]).strip().lower() if row[3] else 'member'
                if home_type not in ['member', 'non_member']:
                    home_type = 'member'
                contact_number = str(row[4]).strip() if row[4] else ''
                address = str(row[5]).strip() if row[5] else ''
                area_name = str(row[6]).strip() if row[6] else ''
                fee_exception_str = str(row[7]).strip().lower() if row[7] else 'no'
                fee_exception = fee_exception_str in ['yes', 'true', '1']
                
                area_obj = None
                if area_name:
                    area_obj, _ = Area.objects.get_or_create(name=area_name)
                
                Home.objects.create(
                    name=name,
                    name_en=name_en,
                    house_name=house_name,
                    home_type=home_type,
                    contact_number=contact_number,
                    address=address,
                    area=area_obj,
                    fee_exception=fee_exception
                )
                count += 1
            
            messages.success(request, f"Successfully imported {count} homes.")
            return redirect('homes:home_list')
        except Exception as e:
            messages.error(request, f"Error processing file: {str(e)}")
            return redirect('homes:home_upload_excel')

    return render(request, 'homes/home_upload_excel.html')

@login_required
def home_download_template(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Homes Template"
    
    # Headers
    headers = ["Name", "Name (English)", "House Name", "Type", "Contact", "Address", "Area", "Fee Exception"]
    ws.append(headers)
    
    # Add a sample row to guide the user
    sample_row = ["John Doe", "John Doe", "Doe Villa", "member", "1234567890", "123 Sample St", "Downtown", "No"]
    ws.append(sample_row)
    
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = adjusted_width

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="homes_upload_template.xlsx"'
    wb.save(response)
    
    return response

@login_required
def area_list(request):
    areas = Area.objects.all().order_by('id')
    return render(request, 'homes/area_list.html', {'areas': areas})

@login_required
def area_create(request):
    if request.method == 'POST':
        form = AreaForm(request.POST)
        if form.is_valid():
            custom_id = form.cleaned_data.get('id')
            if custom_id:
                if Area.objects.filter(id=custom_id).exists():
                    form.add_error('id', 'An area with this ID already exists.')
                    return render(request, 'homes/area_form.html', {'form': form, 'title': 'Create Area'})
            
            area = form.save(commit=False)
            if custom_id:
                area.id = custom_id
            area.save()
            
            messages.success(request, "Area created successfully.")
            return redirect('homes:area_list')
    else:
        form = AreaForm()
    return render(request, 'homes/area_form.html', {'form': form, 'title': 'Create Area'})

@login_required
def area_update(request, pk):
    area = get_object_or_404(Area, pk=pk)
    if request.method == 'POST':
        form = AreaForm(request.POST, instance=area)
        if form.is_valid():
            new_id = form.cleaned_data.get('id')
            old_id = area.id
            
            if new_id and old_id != new_id:
                if Area.objects.filter(id=new_id).exists():
                    form.add_error('id', 'An area with this ID already exists.')
                    return render(request, 'homes/area_form.html', {'form': form, 'title': 'Update Area', 'edit': True})
                
                linked_home_ids = list(area.homes.values_list('id', flat=True))
                Home.objects.filter(id__in=linked_home_ids).update(area=None)
                
                Area.objects.filter(id=old_id).update(
                    id=new_id,
                    name=form.cleaned_data['name'],
                    name_en=form.cleaned_data['name_en']
                )
                
                Home.objects.filter(id__in=linked_home_ids).update(area_id=new_id)
                
                messages.success(request, "Area updated successfully.")
                return redirect('homes:area_list')
            else:
                form.save()
                messages.success(request, "Area updated successfully.")
                return redirect('homes:area_list')
    else:
        form = AreaForm(instance=area)
    return render(request, 'homes/area_form.html', {'form': form, 'title': 'Update Area', 'edit': True})

@login_required
def area_delete(request, pk):
    area = get_object_or_404(Area, pk=pk)
    if request.method == 'POST':
        area.delete()
        messages.success(request, f"Area '{area.name}' deleted successfully.")
        return redirect('homes:area_list')
    return render(request, 'homes/area_confirm_delete.html', {'area': area})

@login_required
@require_POST
def area_create_ajax(request):
    name = request.POST.get('area_name', '').strip()
    name_en = request.POST.get('area_name_en', '').strip()
    
    if not name:
        return JsonResponse({'success': False, 'errors': 'Area name is required.'})
    
    # Check if duplicate exists (case-insensitive)
    existing_area = Area.objects.filter(name__iexact=name).first()
    if existing_area:
        return JsonResponse({
            'success': True,
            'id': existing_area.id,
            'name': existing_area.name,
            'already_exists': True
        })
    
    try:
        area = Area.objects.create(name=name, name_en=name_en)
        return JsonResponse({
            'success': True,
            'id': area.id,
            'name': area.name
        })
    except Exception as e:
        return JsonResponse({'success': False, 'errors': str(e)})

@login_required
def home_list_by_area(request):
    areas = Area.objects.all().order_by('id')
    homes_without_area = Home.objects.filter(area__isnull=True).order_by('id')
    is_print = request.GET.get('print') == 'true'
    total_homes_count = Home.objects.count()
    
    context = {
        'areas': areas,
        'homes_without_area': homes_without_area,
        'is_print': is_print,
        'total_homes_count': total_homes_count,
    }
    
    if is_print:
        return render(request, 'homes/home_list_by_area_print.html', context)
    return render(request, 'homes/home_list_by_area.html', context)
